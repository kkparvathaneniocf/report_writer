import json
import sqlite3
from pathlib import Path

import pytest
from docx import Document

from lab_report.pipeline.docx_utils import TOKEN, replace_text
from lab_report.pipeline.dummy import make_dummy_tr_folder
from lab_report.pipeline.llm import LLMAnswer
from lab_report.pipeline.run import load_spec, run_pipeline

ROOT = Path(__file__).parents[1]
SPEC = ROOT / "config" / "report_spec.yaml"
DNV = ROOT / "data" / "templet" / "DNV report_description for agents AI.docx"


class FakeLLM:
    """Answers from the first snippet containing a keyword; cites that snippet."""

    name = "fake-llm"

    def __init__(self, cite_bogus: bool = False):
        self.cite_bogus = cite_bogus

    def answer(self, system: str, prompt: str) -> str:
        for line in prompt.splitlines():
            if line.startswith("[") and "Unidirectional" in line:
                block_id, text = line[1:].split("] ", 1)
                ids = ["zzzz:p9:l9"] if self.cite_bogus else [block_id]
                return LLMAnswer(value=text, evidence_ids=ids).model_dump_json()
        return LLMAnswer().model_dump_json()


def small_template(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("TR: {{tr_number}} / Product: {{product_description}}")
    doc.add_paragraph("Fabric batch {{fabric_batch}}, catalyst batch {{catalyst_batch}}")
    doc.add_paragraph("Wf mean {{fwf.wf.mean}} (SD {{fwf.wf.sd}}), Rk {{rk_dnv.wf}}")
    table = doc.add_table(rows=2, cols=3)
    for c, h in enumerate(["Specimen", "Wf", "Vf"]):
        table.cell(0, c).text = h
    for c, k in enumerate(["specimen", "wf", "vf"]):
        table.cell(1, c).text = f"{{{{fwf_results[].{k}}}}}"
    doc.save(str(path))
    return path


@pytest.fixture
def tr_folder(tmp_path):
    make_dummy_tr_folder(tmp_path / "TR")
    return tmp_path / "TR"


def slots(prov: dict) -> dict:
    return {s["slot"]: s for s in prov["slots"]}


def test_pipeline_traceability(tmp_path, tr_folder):
    out = tmp_path / "out"
    prov = run_pipeline(tr_folder, out, SPEC, small_template(tmp_path / "t.docx"), FakeLLM())
    s = slots(prov)

    assert {d["doc_type"] for d in prov["documents"]} == {
        "TR_REQUEST",
        "FWF",
        "MACHINE_TENSILE",
        "COA",
        "INFUSION",
    }
    assert s["fwf.wf.mean"]["value"] == "70,18" and s["fwf.wf.mean"]["status"] == "COMPUTED"
    assert len(s["fwf.wf.mean"]["inputs"]) == 5
    assert s["fabric_batch"]["status"] == "CONFLICT"  # planted: COA vs infusion report
    assert s["catalyst_batch"]["status"] == "MISSING"  # planted: absent from infusion report
    assert s["rk_dnv.wf"]["status"] == "NOT_IMPLEMENTED"
    assert s["product_description"]["status"] == "NEEDS_REVIEW"  # LLM output never VERIFIED
    assert s["product_description"]["model"] == "fake-llm"

    # Every cited block exists in the SQLite index, with a crop image on disk.
    with sqlite3.connect(out / "index.sqlite") as con:
        indexed = {row[0] for row in con.execute("SELECT id FROM blocks")}
    for slot in prov["slots"]:
        for e in slot["evidence"]:
            assert e["id"] in indexed
            assert Path(e["crop"]).exists()
        if slot["status"] == "VERIFIED":
            assert slot["evidence"], slot["slot"]

    doc = Document(str(out / "report.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Wf mean 70,18 (SD 0,43)" in text
    assert "catalyst batch [MISSING]" in text
    rows = [[c.text for c in r.cells] for r in doc.tables[0].rows]
    assert len(rows) == 6 and rows[1] == ["TR26_C0143_01_1", "69,64", "50,18"]
    assert len(list(doc.comments)) > 5
    assert not prov["unfilled_tokens"]
    assert json.loads((out / "provenance.json").read_text())["run"]["llm"] == "fake-llm"


def test_llm_bogus_citation_is_flagged(tmp_path, tr_folder):
    prov = run_pipeline(
        tr_folder,
        tmp_path / "o",
        SPEC,
        small_template(tmp_path / "t.docx"),
        FakeLLM(cite_bogus=True),
    )
    desc = slots(prov)["product_description"]
    assert desc["status"] == "NEEDS_REVIEW"
    assert desc["evidence_ids"] == [] and "invalid or no citations" in desc["reason"]


def test_without_llm_fields_are_missing(tmp_path, tr_folder):
    prov = run_pipeline(tr_folder, tmp_path / "o", SPEC, small_template(tmp_path / "t.docx"))
    assert slots(prov)["product_description"]["status"] == "MISSING"


def test_replace_text_across_runs_keeps_formatting():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Density: ")
    bold = p.add_run("x")
    bold.bold = True
    p.add_run(" g/cm3")
    run = replace_text(p, "x g", "{{rho}} g")
    assert p.text == "Density: {{rho}} g/cm3"
    assert run is not None and run.bold


def test_spec_rejects_boolean_yaml_keys(tmp_path):
    bad = tmp_path / "spec.yaml"
    bad.write_text("tables:\n  t:\n    derive:\n      no: {from: a, regex: x}\n")
    with pytest.raises(ValueError, match="not a string"):
        load_spec(bad)


@pytest.mark.skipif(not DNV.exists(), reason="DNV template not available (not in git)")
def test_full_dnv_template(tmp_path, tr_folder):
    from lab_report.pipeline.template_builder import build_template

    template = build_template(DNV, tmp_path / "dnv.docx")
    assert not list(Document(str(template)).comments)  # authoring notes stripped
    prov = run_pipeline(tr_folder, tmp_path / "o", SPEC, template, FakeLLM())
    assert not prov["unfilled_tokens"]
    doc = Document(str(tmp_path / "o" / "report.docx"))
    for p in doc.paragraphs:
        assert not TOKEN.search(p.text)


class FakeOllamaClient:
    """Stands in for ollama.Client. mode: 'thinking' = always empty content with thinking
    text (the reported Gemma 4 failure); 'schema-empty' = empty only with a JSON schema."""

    mode = "thinking"
    calls: list = []

    def __init__(self, host):
        pass

    def list(self):
        from types import SimpleNamespace

        return SimpleNamespace(models=[SimpleNamespace(model="gemma4:26b", digest="08ae7ec1744b")])

    def chat(self, model, messages, format, think, options):
        from types import SimpleNamespace

        FakeOllamaClient.calls.append({"format": format, "think": think, **options})
        empty = self.mode == "thinking" or isinstance(format, dict)
        content = "" if empty else LLMAnswer(value="UD laminate").model_dump_json()
        message = SimpleNamespace(content=content, thinking="hmm " * 400 if empty else None)
        return SimpleNamespace(message=message, done_reason="length", eval_count=2048)


def test_ollama_empty_answer_does_not_stop_report(tmp_path, tr_folder, monkeypatch):
    import ollama

    from lab_report.pipeline.llm import OllamaClient

    monkeypatch.setattr(ollama, "Client", FakeOllamaClient)
    FakeOllamaClient.mode, FakeOllamaClient.calls = "thinking", []
    llm = OllamaClient("gemma4:26b", "http://x")
    prov = run_pipeline(tr_folder, tmp_path / "o", SPEC, small_template(tmp_path / "t.docx"), llm)
    desc = slots(prov)["product_description"]
    assert desc["status"] == "MISSING"
    assert "empty response" in desc["reason"] and "thinking_chars=1600" in desc["reason"]
    assert (tmp_path / "o" / "report.docx").exists()
    assert all(c["think"] is False for c in FakeOllamaClient.calls)
    assert FakeOllamaClient.calls[1]["format"] == "json"  # retried without the schema


def test_ollama_schema_failure_falls_back_to_json_mode(tmp_path, tr_folder, monkeypatch):
    import ollama

    from lab_report.pipeline.llm import OllamaClient

    monkeypatch.setattr(ollama, "Client", FakeOllamaClient)
    FakeOllamaClient.mode, FakeOllamaClient.calls = "schema-empty", []
    llm = OllamaClient("gemma4:26b", "http://x")
    prov = run_pipeline(tr_folder, tmp_path / "o", SPEC, small_template(tmp_path / "t.docx"), llm)
    desc = slots(prov)["product_description"]
    assert desc["value"] == "UD laminate" and desc["status"] == "NEEDS_REVIEW"
    assert desc["model"] == "ollama:gemma4:26b@08ae7ec1744b"
