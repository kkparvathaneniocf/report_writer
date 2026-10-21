from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

from docx import Document
import fitz
import pandas as pd


@dataclass(frozen=True)
class Fragment:
    source_file: str
    source_sha256: str
    location: str
    text: str


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_file(path: Path) -> list[Fragment]:
    digest = sha256(path)
    parts: list[Fragment] = []
    suffix = path.suffix.lower()
    if suffix == ".docx":
        doc = Document(path)
        for i, p in enumerate(doc.paragraphs, 1):
            if p.text.strip():
                parts.append(Fragment(path.name, digest, f"paragraph:{i}", p.text.strip()))
        for t, table in enumerate(doc.tables, 1):
            for r, row in enumerate(table.rows, 1):
                value = " | ".join(cell.text.replace("\n", " ").strip() for cell in row.cells)
                if value.strip(" |"):
                    parts.append(Fragment(path.name, digest, f"table:{t}/row:{r}", value))
    elif suffix == ".pdf":
        with fitz.open(path) as pdf:
            for i, page in enumerate(pdf, 1):
                text = page.get_text("text").strip()
                if text:
                    parts.append(Fragment(path.name, digest, f"page:{i}", text))
    elif suffix in {".csv", ".xlsx"}:
        if suffix == ".csv":
            df = pd.read_csv(path, sep=None, engine="python", dtype=str).fillna("")
            sheets = {"csv": df}
        else:
            sheets = pd.read_excel(path, sheet_name=None, dtype=str)
        for sheet, frame in sheets.items():
            frame = frame.fillna("")
            for i, (_, row) in enumerate(frame.iterrows(), 2):
                value = " | ".join(f"{column}={row[column]}" for column in frame.columns)
                parts.append(Fragment(path.name, digest, f"sheet:{sheet}/row:{i}", value))
    elif suffix in {".txt", ".md"}:
        for i, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            if line.strip():
                parts.append(Fragment(path.name, digest, f"line:{i}", line.strip()))
    return parts


def ingest_folder(folder: Path) -> list[Fragment]:
    if not folder.exists():
        raise FileNotFoundError(f"Input folder does not exist: {folder}")
    out: list[Fragment] = []
    for path in sorted(folder.rglob("*")):
        if path.is_file() and not path.name.startswith("~$"):
            out.extend(read_file(path))
    return out
