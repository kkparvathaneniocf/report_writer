"""Turn the DNV report description template into a tokenized POC template.

Placeholders like "xyz", "Name" or "X samples" become {{slot}} tokens; repeated table rows become
one template row with {{group[].column}} tokens. Fixed and manual (yellow) content is untouched.
The mapping is specific to "DNV report_description for agents AI.docx" (table order as of
2026-10); if the template changes, adjust the indices below and rerun.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from lab_report.pipeline.docx_utils import (
    delete_paragraph,
    delete_row,
    replace_text,
    set_cell_text,
)


def _para(paragraphs: list[Paragraph], contains: str, nth: int = 0) -> Paragraph:
    matches = [p for p in paragraphs if contains in p.text]
    if len(matches) <= nth:
        raise ValueError(f"template paragraph containing {contains!r} (#{nth}) not found")
    return matches[nth]


def _sub(paragraphs: list[Paragraph], contains: str, old: str, new: str, nth: int = 0) -> None:
    if replace_text(_para(paragraphs, contains, nth), old, new) is None:
        raise ValueError(f"{old!r} not found in paragraph containing {contains!r}")


def _row_tokens(table, row_index: int, tokens: dict[int, str]) -> None:
    for col, token in tokens.items():
        set_cell_text(table.cell(row_index, col), token)


def _insert_tensile_results_table(doc: DocumentType, anchor: Paragraph) -> None:
    headers = [
        "Specimen N°",
        "Thickness [mm]",
        "Stress at break [MPa]",
        "Modulus [GPa]",
        "Lin. est. strain [%]",
    ]
    cols = ["specimen", "thickness", "stress", "modulus", "strain"]
    table = doc.add_table(rows=5, cols=len(headers))
    table.style = doc.tables[1].style  # "Table Grid", as used across the template
    for c, h in enumerate(headers):
        run = table.cell(0, c).paragraphs[0].add_run(h)
        run.bold = True
    for c, key in enumerate(cols):
        table.cell(1, c).paragraphs[0].add_run(f"{{{{tensile_results[].{key}}}}}")
    for r, (label, op) in enumerate((("Average", "mean"), ("SD", "sd"), ("CV [%]", "cv")), 2):
        table.cell(r, 0).paragraphs[0].add_run(label).bold = True
        for c, key in enumerate(cols[1:], 1):
            table.cell(r, c).paragraphs[0].add_run(f"{{{{tensile.{key}.{op}}}}}")
    for row in table.rows:
        for cell in row.cells:
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
    anchor._p.addprevious(table._tbl)
    delete_paragraph(anchor)


def _strip_comments(doc: DocumentType) -> None:
    """Remove the template author's notes ("Source = page 1", ...); the pipeline adds its own
    provenance comments and the two must not be confused in the generated report."""
    for tag in ("w:commentRangeStart", "w:commentRangeEnd"):
        for el in list(doc.element.body.iter(qn(tag))):
            el.getparent().remove(el)
    for ref in list(doc.element.body.iter(qn("w:commentReference"))):
        run = ref.getparent()
        run.getparent().remove(run)
    comments = doc.part._comments_part.element
    for c in list(comments):
        comments.remove(c)


def build_template(source: Path, dest: Path) -> Path:
    doc = Document(str(source))
    _strip_comments(doc)
    t = doc.tables
    body = doc.paragraphs

    # Cover page (table 0) and general information (table 1)
    cover = t[0].cell(0, 0).paragraphs
    p = _para(cover, "Product name & brief description")
    replace_text(p, "“Product name & brief description “", "{{product_description}}")
    _sub(cover, "Test(s) description", "Test(s) description", "{{tests_description}}")
    _sub(cover, "TR # associated", "TR # associated", "{{tr_number}}")
    tasks = t[1].cell(2, 1).paragraphs
    replace_text(tasks[0], "Product name & brief description", "{{product_name}}")
    replace_text(tasks[1], "TEST 1 + associated STANDARD REF", "{{tasks_tests}}")
    for extra in tasks[2:]:
        delete_paragraph(extra)

    # Scope
    scope = _para(body, "Context and summary")
    replace_text(scope, scope.text, "{{scope_summary}}")

    # Results summary (table 3): FWF columns 1-3, tensile 0° columns 4-6
    summary = t[3]
    for col, label in (
        (4, "Tensile 0°\nσ break (MPa)"),
        (5, "Tensile 0°\nE (GPa)"),
        (6, "Tensile 0°\nε lin (%)"),
    ):
        set_cell_text(summary.cell(0, col), label).bold = True
    for row, op in enumerate(("mean", "sd", "cv", "max", "min"), 1):
        _row_tokens(
            summary,
            row,
            {
                1: f"{{{{fwf.wf.{op}}}}}",
                2: f"{{{{fwf.vf.{op}}}}}",
                3: f"{{{{fwf.porosity.{op}}}}}",
                4: f"{{{{tensile.stress.{op}}}}}",
                5: f"{{{{tensile.modulus.{op}}}}}",
                6: f"{{{{tensile.strain.{op}}}}}",
            },
        )
    for row, rk in ((6, "rk_dnv"), (7, "rk_gl")):
        _row_tokens(
            summary,
            row,
            {
                1: f"{{{{{rk}.wf}}}}",
                2: f"{{{{{rk}.vf}}}}",
                3: f"{{{{{rk}.porosity}}}}",
                4: f"{{{{{rk}.tensile}}}}",
            },
        )
    _sub(body, "Density of resin: x", "resin: x", "resin: {{resin_density}}")
    _sub(body, "density of fiber: y", "fiber: y", "fiber: {{fiber_density}}")

    # Standards (table 4) and internal procedures (table 5): one template row each
    _row_tokens(
        t[4], 1, {0: "{{standards[].number}}", 1: "{{standards[].title}}", 2: "{{standards[].dnv}}"}
    )
    _row_tokens(t[5], 1, {0: "{{procedures[].number}}", 1: "{{procedures[].title}}"})
    for table in (t[4], t[5]):
        for row in list(table.rows)[2:]:
            delete_row(row)

    # Materials (tables 6, 7) and manufacturing parameters (table 8)
    _row_tokens(
        t[6],
        1,
        dict(
            enumerate(
                [
                    "{{fabric_description}}",
                    "{{fabric_supplier}}",
                    "{{fabric_production_date}}",
                    "{{fabric_code}}",
                    "{{fabric_batch}}",
                    "{{fabric_roll}}",
                    "{{fabric_plant}}",
                ]
            )
        ),
    )
    _row_tokens(t[7], 1, {1: "{{resin_name}}", 2: "{{resin_supplier}}", 3: "{{resin_batch}}"})
    _row_tokens(
        t[7], 2, {1: "{{catalyst_name}}", 2: "{{catalyst_supplier}}", 3: "{{catalyst_batch}}"}
    )
    _row_tokens(
        t[8],
        1,
        {
            0: "{{laminate_id}}",
            2: "{{depression}}",
            3: "{{infusion_temp}}",
            4: "{{curing_temp}}",
            5: "{{postcure}}",
            6: "{{sop}}",
        },
    )

    # Specimen identification (table 10): one row per test
    _row_tokens(
        t[10],
        3,
        {
            0: "{{specimen_ids[].tr}}",
            1: "{{specimen_ids[].refs}}",
            2: "{{specimen_ids[].test}}",
            3: "{{specimen_ids[].direction}}",
            4: "{{specimen_ids[].count}}",
        },
    )

    # Fiber content section
    _sub(body, "Technician: P. BERTELONE", "P. BERTELONE", "{{fwf_technician}}")
    _sub(body, "Number of specimens: 5", "specimens: 5", "specimens: {{fwf.count}}")
    _row_tokens(
        t[11],
        1,
        {
            0: "{{fwf_equipment[].number}}",
            1: "{{fwf_equipment[].name}}",
            2: "{{fwf_equipment[].description}}",
            3: "{{fwf_equipment[].calibration}}",
        },
    )
    delete_row(t[11].rows[2])
    fwf = t[12]
    _row_tokens(
        fwf,
        1,
        {
            0: "{{fwf_results[].ref}}",
            1: "{{fwf_results[].num}}",
            2: "{{fwf_results[].wf}}",
            3: "{{fwf_results[].vf}}",
            4: "{{fwf_results[].density}}",
            5: "{{fwf_results[].porosity}}",
        },
    )
    for row in list(fwf.rows)[2:6]:
        delete_row(row)
    for row, op in ((2, "mean"), (3, "sd"), (4, "cv")):
        _row_tokens(
            fwf,
            row,
            {
                c: f"{{{{fwf.{k}.{op}}}}}"
                for c, k in ((2, "wf"), (3, "vf"), (4, "density"), (5, "porosity"))
            },
        )
    _sub(body, "Table 9 - TR26_C0143_01", "TR26_C0143_01", "{{fwf_results[0].ref}}")
    _sub(body, "Density of resin: X", "resin: X", "resin: {{resin_density}} ")
    _sub(body, "fiber: Y g/cm3", "fiber: Y", "fiber: {{fiber_density}}")

    # Tensile 0° section ("Test Name A")
    _sub(
        body,
        "Test Name A according to",
        "Test Name A according to Standard reference",
        "{{tensile_title}}",
    )
    _sub(body, "Technician: Name", "Name", "{{tensile_technician}}")
    _row_tokens(
        t[14],
        1,
        {
            0: "{{tensile_equipment[].number}}",
            1: "{{tensile_equipment[].name}}",
            2: "{{tensile_equipment[].description}}",
        },
    )
    for row in list(t[14].rows)[2:]:
        delete_row(row)
    _sub(body, "Number of specimens: X samples", "X samples", "{{tensile.count}} samples")
    _sub(body, "Test Name A - Specimen geometry", "Test Name A", "{{tensile_test}}")
    set_cell_text(t[15].cell(2, 1), "{{tensile.thickness.mean}}")
    _sub(body, "Prestress: X MPa", "X MPa", "{{tensile_prestress}}")
    _sub(body, "Test speed = Y mm/min", "Y mm/min", "{{tensile_speed}}")
    _insert_tensile_results_table(doc, _para(body, "Table of results from the CSV file"))

    # Appendices: page snapshots, curves and specimen pictures ({{image:<name>}} tokens)
    _sub(
        body,
        "Picture / screenshot of the COA",
        "Picture / screenshot of the COA from the TR folder",
        "{{image:coa_snapshot}}",
    )
    _sub(body, "Infusion report – Laminate #", "260612VB-A", "{{laminate_id}}")
    _sub(
        body,
        "Picture / Screeshot or from TR folder",
        "Picture / Screeshot or from TR folder or directly from WaveFront ?",
        "{{image:infusion_snapshot}}",
    )
    for _ in range(2):  # "Test report" and "Pictures" appendix headings
        _sub(
            body,
            "TRYY_CXXXX– Test name",
            "TRYY_CXXXX– Test name",
            "{{tr_number}} – {{tensile_test}}",
        )
    _sub(
        body,
        "Machine report + stress/strain curves",
        "Machine report + stress/strain curves",
        "{{image:tensile_report}}{{image:tensile_curves}}",
    )
    # The template repeats "Picture before test"; the second one is taken as "after test".
    _sub(body, "Picture before test", "Picture before test", "{{image:pictures_before}}")
    _sub(body, "Picture before test", "Picture before test", "{{image:pictures_after}}")

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return dest
