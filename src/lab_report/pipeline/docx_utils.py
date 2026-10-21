"""Formatting-preserving text replacement for python-docx documents."""

from __future__ import annotations

import re
from collections.abc import Iterator
from copy import deepcopy

from docx.document import Document
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph
from docx.text.run import Run

TOKEN = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def replace_text(paragraph: Paragraph, old: str, new: str) -> Run | None:
    """Replace the first `old` in `paragraph`, even when it spans several runs.

    The new text takes the formatting of the run where `old` starts; text before and after
    keeps its own runs. Returns the run holding `new` (to attach comments/highlights).
    """
    runs = paragraph.runs
    full = "".join(r.text for r in runs)
    start = full.find(old)
    if start < 0 or not old:
        return None
    end = start + len(old)
    bounds, pos = [], 0
    for r in runs:
        bounds.append((pos, pos + len(r.text)))
        pos += len(r.text)
    s = next(i for i, (a, b) in enumerate(bounds) if a <= start < b)
    e = next(i for i, (a, b) in enumerate(bounds) if a < end <= b)
    before = runs[s].text[: start - bounds[s][0]]
    after = runs[e].text[end - bounds[e][0] :]

    new_el = deepcopy(runs[s]._r)
    runs[s]._r.addnext(new_el)
    new_run = Run(new_el, paragraph)
    new_run.text = new
    if s == e:
        runs[s].text = before
        if after:
            tail = deepcopy(runs[s]._r)
            new_el.addnext(tail)
            Run(tail, paragraph).text = after
    else:
        runs[s].text = before
        for r in runs[s + 1 : e]:
            r.text = ""
        runs[e].text = after
    return new_run


def iter_tables(container) -> Iterator[Table]:
    for table in container.tables:
        yield table
        for row in table.rows:
            for cell in row.cells:
                yield from iter_tables(cell)


def iter_paragraphs(doc: Document) -> Iterator[Paragraph]:
    """Body, table (incl. nested) and header/footer paragraphs, each once.

    `seen` holds the lxml elements themselves: keeping them alive keeps their proxies stable
    (id() of a discarded proxy can be reused, which silently skips paragraphs).
    """
    seen: set = set()

    def from_container(container) -> Iterator[Paragraph]:
        for p in container.paragraphs:
            if p._p not in seen:
                seen.add(p._p)
                yield p
        for table in iter_tables(container):
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p._p not in seen:
                            seen.add(p._p)
                            yield p

    yield from from_container(doc)
    for section in doc.sections:
        for part in (section.header, section.footer):
            if not part.is_linked_to_previous:
                yield from from_container(part)


def row_paragraphs(row: _Row) -> list[Paragraph]:
    out: list[Paragraph] = []
    seen: set = set()
    for cell in row.cells:
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        out.extend(cell.paragraphs)
    return out


def clone_row_after(row: _Row, anchor: _Row | None = None) -> _Row:
    new_tr = deepcopy(row._tr)
    (anchor or row)._tr.addnext(new_tr)
    return _Row(new_tr, row._parent)  # type: ignore[arg-type]


def delete_row(row: _Row) -> None:
    row._tr.getparent().remove(row._tr)


def delete_paragraph(paragraph: Paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def set_cell_text(cell, text: str) -> Run:
    """Put `text` into a cell's first paragraph, reusing its first run's formatting."""
    paragraphs = cell.paragraphs
    for extra in paragraphs[1:]:
        delete_paragraph(extra)
    p = paragraphs[0]
    runs = p.runs
    if not runs:
        return p.add_run(text)
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""
    return runs[0]
