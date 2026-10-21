from __future__ import annotations

from pathlib import Path

from docx import Document

from lab_report.models import AgentOutput


def write_report(
    outputs: list[AgentOutput], destination: Path, template: Path | None = None
) -> None:
    """Template tokens: {{field_name}} in body paragraphs or table cells.

    Token replacement runs within each paragraph; token can span multiple runs.
    Replacing paragraph text can lose inline styling for paragraphs containing tokens.
    Do not use this on the production report without testing on a copy.
    """
    doc = Document(str(template)) if template else Document()
    fields = {f.field_name: f.value for item in outputs for f in item.fields if f.value is not None}
    if template:
        paragraphs = list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        for paragraph in paragraphs:
            original = paragraph.text
            updated = original
            for name, value in fields.items():
                updated = updated.replace("{{" + name + "}}", value)
            if updated != original:
                paragraph.text = updated
    else:
        doc.add_heading("Laboratory Report — Draft", 0)
        for item in outputs:
            doc.add_heading(f"Section {item.section_id}", level=1)
            for field in item.fields:
                doc.add_paragraph(
                    f"{field.field_name}: {field.value or '[MISSING]'} [{field.status}] "
                    f"Evidence: {', '.join(field.evidence_ids)}"
                )
            if item.narrative:
                doc.add_paragraph(item.narrative)
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destination))
